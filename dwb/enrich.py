"""Enrich stored Orders with the Charges in a History export.

The API reports an Order's Final Price as a single number; the itemisation —
base price, waiting time, extra stops — exists only in a History export's
PriceBreakdown column. This run parses that column (dwb/price_breakdown.py)
and stores the result: one order_price_breakdowns row per Order carrying the
string as exported, and one order_charges row per Charge.

Export rows match stored Orders on the export's ID column, which is the
dispatch system's own record id — never the Order Number, which Digital
Waybill reissues; matching on a reissued number would attach one Order's
Charges to its twin. An ID with no stored Order is counted and reported, not
treated as an error: the store is known to be missing Orders that only
exports still name.

Re-running is safe: each Order's breakdown row is upserted and its Charges
are replaced wholesale, so a corrected export supersedes an earlier run
instead of accumulating beside it.

Run from the repository root, where the wrapper lives:

    uv run python enrich_price_breakdown.py            # input/History-2025-04-12.xlsx
    uv run python enrich_price_breakdown.py --input input/History-2026-01-31.xlsx
    uv run python enrich_price_breakdown.py --batch-size 1000
    uv run python enrich_price_breakdown.py --source another_tms
    uv run python enrich_price_breakdown.py --no-progress

To check an export parses before letting it near the store, parse_price_breakdown.py
validates the same column and writes nothing.

Exit status: 0 on success, 2 if the export file is missing or the store is
unreachable. A row whose charges miss FinalPrice is stored and flagged, not a
failure.
"""

import argparse
import os
import sys

import openpyxl
from psycopg.types.json import Jsonb
from tqdm import tqdm

from dwb import db
from dwb.config import load_dotenv
from dwb.ingest import DEFAULT_SOURCE
from dwb.price_breakdown import parse

DEFAULT_INPUT = "input/History-2025-04-12.xlsx"
DEFAULT_BATCH_SIZE = 500

# The export columns enrichment needs; everything else in the row is ignored.
EXPORT_COLUMNS = ("ID", "PriceBreakdown", "FinalPrice")


def read_export(path):
    """Yield (dispatch_record_id, text, breakdown, final_price) per data row."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = {name: i for i, name in enumerate(next(rows))}
        for column in EXPORT_COLUMNS:
            if column not in header:
                raise SystemExit(f"{path}: missing column {column!r}")
        for row in rows:
            text = row[header["PriceBreakdown"]]
            yield (row[header["ID"]], text, parse(text),
                   row[header["FinalPrice"]])
    finally:
        workbook.close()


class Summary:
    """What an enrichment run did, honestly."""

    def __init__(self):
        self.rows_read = 0
        self.orders_enriched = 0
        self.charges_written = 0
        self.sum_mismatches = 0    # stored anyway, flagged in the row
        self.unmatched = []        # export IDs naming no stored Order

    def lines(self):
        """The summary as printable lines, unmatched export IDs listed to a cap."""
        yield f"Rows read:         {self.rows_read}"
        yield f"Orders enriched:   {self.orders_enriched}"
        yield f"Charges written:   {self.charges_written}"
        yield (f"Sum != FinalPrice: {self.sum_mismatches}"
               " (stored, matches_final_price = false)")
        yield f"No stored Order:   {len(self.unmatched)}"
        for record_id in self.unmatched[:20]:
            yield f"  export ID {record_id}"
        if len(self.unmatched) > 20:
            yield f"  … and {len(self.unmatched) - 20} more"


UPSERT_BREAKDOWN_SQL = """
    insert into order_price_breakdowns
        (order_id, source_key, breakdown, source_file,
         export_final_price, matches_final_price)
    values (%s, %s, %s, %s, %s, %s)
    on conflict (order_id) do update set
        breakdown = excluded.breakdown,
        source_file = excluded.source_file,
        export_final_price = excluded.export_final_price,
        matches_final_price = excluded.matches_final_price,
        enriched_at = now()
"""

INSERT_CHARGE_SQL = """
    insert into order_charges
        (order_id, charge_position, description, quantity, rate,
         pricing_code, refs)
    values (%s, %s, %s, %s, %s, %s, %s)
"""


def _flush(conn, batch, source_key, source_file, summary):
    """Write one batch of export rows and commit."""
    if not batch:
        return
    with conn.cursor() as cur:
        cur.execute(
            "select dispatch_record_id, order_id from orders "
            "where source_key = %s and dispatch_record_id = any(%s)",
            (source_key, [record_id for record_id, _, _, _ in batch]))
        order_ids = dict(cur.fetchall())

        breakdown_rows, charge_rows, matched = [], [], []
        for record_id, text, breakdown, final_price in batch:
            order_id = order_ids.get(record_id)
            if order_id is None:
                summary.unmatched.append(record_id)
                continue
            matched.append(order_id)
            matches = breakdown.matches(final_price)
            if not matches:
                summary.sum_mismatches += 1
            breakdown_rows.append(
                (order_id, source_key, "" if text is None else str(text),
                 source_file, final_price, matches))
            for position, charge in enumerate(breakdown.charges, start=1):
                charge_rows.append(
                    (order_id, position, charge.description, charge.quantity,
                     charge.rate, charge.code, Jsonb(list(charge.refs))))

        if matched:
            cur.executemany(UPSERT_BREAKDOWN_SQL, breakdown_rows)
            cur.execute("delete from order_charges where order_id = any(%s)",
                        (matched,))
            if charge_rows:
                cur.executemany(INSERT_CHARGE_SQL, charge_rows)
            summary.orders_enriched += len(matched)
            summary.charges_written += len(charge_rows)
    conn.commit()


def enrich(conn, input_path, source_key=DEFAULT_SOURCE,
           batch_size=DEFAULT_BATCH_SIZE, show_progress=False):
    """Enrich from every row of the export at input_path. Returns a Summary."""
    summary = Summary()
    source_file = os.path.basename(input_path)
    bar = tqdm(unit="row", desc="Enriching", disable=not show_progress)
    batch = []
    try:
        for record in read_export(input_path):
            summary.rows_read += 1
            batch.append(record)
            bar.update(1)
            if len(batch) >= batch_size:
                _flush(conn, batch, source_key, source_file, summary)
                batch = []
        _flush(conn, batch, source_key, source_file, summary)
    finally:
        bar.close()
    return summary


def main(argv=None):
    """Run an enrichment from the command line. 0 on success, 2 if it could not start.

    Rows whose charges do not sum to FinalPrice are stored and counted, not
    rejected, so a sum mismatch is reported without failing the run.
    """
    load_dotenv()
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", default=DEFAULT_INPUT,
                        help=f"Path to the History .xlsx export (default {DEFAULT_INPUT})")
    parser.add_argument("--dsn", default=None,
                        help="Postgres connection string. Env: DWB_DSN")
    parser.add_argument("--source", default=DEFAULT_SOURCE,
                        help="Which dispatch system the export came from "
                             f"(default {DEFAULT_SOURCE})")
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE,
                        help=f"Export rows per transaction (default {DEFAULT_BATCH_SIZE})")
    parser.add_argument("--progress", dest="progress", action="store_true",
                        default=sys.stderr.isatty(),
                        help="Show a progress bar (default: only on a terminal).")
    parser.add_argument("--no-progress", dest="progress", action="store_false",
                        help="Never show a progress bar.")
    args = parser.parse_args(argv)

    if not os.path.exists(args.input):
        print(f"No such file: {args.input}", file=sys.stderr)
        return 2

    conn = db.connect_or_exit(args.dsn)
    if conn is None:
        return 2

    try:
        with conn:
            summary = enrich(conn, args.input, source_key=args.source,
                             batch_size=args.batch_size,
                             show_progress=args.progress)
    finally:
        conn.close()

    for line in summary.lines():
        print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
